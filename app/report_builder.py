"""Multi-format report builder: CSV, Excel (.xlsx), PDF."""
import csv
import io
from datetime import date, datetime
from typing import Any

from fastapi.responses import StreamingResponse


def _fmt_date(d: date | datetime | str | None) -> str:
    if d is None:
        return ""
    if isinstance(d, str):
        return d[:10]
    return d.strftime("%Y-%m-%d")


def build_csv(rows: list[dict], columns: list[str], headers: list[str] | None = None) -> io.BytesIO:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers or columns)
    for row in rows:
        writer.writerow([row.get(c, "") for c in columns])
    out = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
    out.seek(0)
    return out


def build_excel(
    sheets: dict[str, tuple[list[dict], list[str], list[str] | None]],
    title: str = "Report",
    date_from: str = "",
    date_to: str = "",
) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    title_font = Font(bold=True, size=12, color="0F172A")
    meta_font = Font(italic=True, size=9, color="64748B")

    for sheet_name, (rows, columns, headers) in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 1))
        ws.cell(1, 1, title).font = title_font

        period_str = ""
        if date_from and date_to:
            period_str = f"Period: {date_from} to {date_to}"
        elif date_from:
            period_str = f"From: {date_from}"
        elif date_to:
            period_str = f"To: {date_to}"
        if period_str:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(columns), 1))
            ws.cell(2, 1, period_str).font = meta_font

        start_row = 4
        display_headers = headers or columns
        for col_idx, h in enumerate(display_headers, 1):
            cell = ws.cell(start_row, col_idx, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for row_idx, row in enumerate(rows, start_row + 1):
            for col_idx, col in enumerate(columns, 1):
                val = row.get(col, "")
                if isinstance(val, (date, datetime)):
                    val = _fmt_date(val)
                cell = ws.cell(row_idx, col_idx, val)
                cell.border = thin_border

        from openpyxl.utils import get_column_letter
        for col_idx, col in enumerate(columns, 1):
            max_len = len(display_headers[col_idx - 1]) if col_idx <= len(display_headers) else 10
            for row in rows[:100]:
                v = str(row.get(col, ""))
                max_len = max(max_len, min(len(v), 50))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_pdf(
    rows: list[dict],
    columns: list[str],
    headers: list[str] | None = None,
    title: str = "Report",
    date_from: str = "",
    date_to: str = "",
) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    page_size = landscape(A4) if len(columns) > 5 else A4
    doc = SimpleDocTemplate(buf, pagesize=page_size, leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    elements = []

    cell_style = ParagraphStyle("cell", parent=styles["Normal"], fontSize=6.5,
                                leading=8, alignment=TA_LEFT)
    cell_center = ParagraphStyle("cellc", parent=cell_style, alignment=TA_CENTER)
    header_style = ParagraphStyle("hdr", parent=styles["Normal"], fontSize=7,
                                  leading=9, textColor=colors.white,
                                  fontName="Helvetica-Bold", alignment=TA_CENTER)

    elements.append(Paragraph(title, styles["Title"]))
    if date_from or date_to:
        period = f"Period: {date_from or '...'} to {date_to or 'now'}"
        elements.append(Paragraph(period, styles["Normal"]))
    elements.append(Spacer(1, 4 * mm))

    display_headers = headers or columns
    table_data = [[Paragraph(h, header_style) for h in display_headers]]
    for row in rows:
        cells = []
        for i, c in enumerate(columns):
            val = str(row.get(c, ""))
            style = cell_style if i == 0 else cell_center
            cells.append(Paragraph(val, style))
        table_data.append(cells)

    col_count = len(columns)
    available_width = page_size[0] - 20 * mm
    first_col_ratio = 3.0
    other_ratio = 1.0
    total_ratio = first_col_ratio + other_ratio * (col_count - 1)
    col_widths = [available_width * first_col_ratio / total_ratio]
    for _ in range(col_count - 1):
        col_widths.append(available_width * other_ratio / total_ratio)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)

    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(f"Generated: {generated} | GEO Agent - GrabOn", styles["Normal"]))

    doc.build(elements)
    buf.seek(0)
    return buf


def make_response(
    buf: io.BytesIO,
    fmt: str,
    filename: str,
) -> StreamingResponse:
    media_types = {
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pdf": "application/pdf",
    }
    return StreamingResponse(
        buf,
        media_type=media_types.get(fmt, "application/octet-stream"),
        headers={"Content-Disposition": f'attachment; filename="{filename}.{fmt}"'},
    )
